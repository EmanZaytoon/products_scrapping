import datetime
from time import sleep
from PyQt5.QtWidgets import QApplication, QWidget
from PyQt5.QtGui import *
from PyQt5.QtCore import *
from PyQt5.QtWidgets import *
from PyQt5.uic import loadUi
import sys
import os
from openpyxl import Workbook , load_workbook
from openpyxl.styles import NamedStyle, Font, Border, Side, Alignment
from openpyxl.utils import get_column_letter
import pandas as pd
import random
from bs4 import BeautifulSoup
from seleniumbase import SB
import random

class scrapping_window (QMainWindow):
    def __init__(self):
        super(scrapping_window,self).__init__()
        loadUi("scrapping.ui",self)
        self.lineEdit.setFocus()
        self.url=''
        self.html_source=''
        self.buttons()
        self.search=True
    def keyPressEvent(self, event):
        if event.key() in (Qt.Key_Return, Qt.Key_Enter):
            self.search()
        else:
         super().keyPressEvent(event)
    def buttons(self):
        self.pushButton.clicked.connect(self.search)
        self.pushButton_3.clicked.connect(self.search_all)
        self.pushButton.setDefault(True)
        self.pushButton_2.clicked.connect(self.open_excel)
        self.pushButton_4.clicked.connect(self.stop_searching)
    def search(self):
        '''self.tableWidget.clearContents()
        self.tableWidget.setHorizontalHeaderLabels([])
        QApplication.processEvents()'''
        self.search_word=self.lineEdit.text()
        self.site_name=self.comboBox.currentText()
        sites={'Amazon':'https://www.amazon.eg/ref=nav_logo','Jumia':'https://www.jumia.com.eg/ar/','Noon':f'https://www.noon.com/egypt-en/search/?q={self.search_word}'}
        url=f'{sites[self.site_name]}'
        with SB(uc=True, test=True, locale="en") as sb:
            sb.sleep(random.randint(1, 5))
            sb.maximize()
            sb.sleep(random.randint(1, 5))
            sb.activate_cdp_mode(url)         
            sb.sleep(random.randint(5, 10))
            try:
                sb.solve_captcha()
                sb.sleep(random.randint(5, 7))
            except:
                print('nocaptcha')  
            if self.site_name =='Amazon':
                sb.type('input[id="twotabsearchtextbox"]',self.search_word)
                self.search_button='input[id="nav-search-submit-button"]'
                sb.click(self.search_button)
                sb.sleep(random.randint(1,5))
                self.html_source = sb.get_page_source()
                self.amazone()
                QMessageBox.about(self,'Finish', 'Data scrapped successfully!')  
            elif self.site_name =='Jumia':
                sb.click('button.cls')
                sb.sleep(random.randint(1, 3))
                sb.click('Button.-df')
                sb.type('input[name="q"]', self.search_word ,timeout=5)
                sb.sleep(random.randint(5, 10))
                sb.click('Button._search')
                sb.sleep(random.randint(5, 10))
                self.html_source = sb.get_page_source() 
                self.jumia()
            elif self.site_name =='Noon': 
                self.html_source = sb.get_page_source()
                self.noon()     
    def search_all(self):
        self.search_word=self.lineEdit.text()
        self.site_name=self.comboBox.currentText()
        sites={'Amazon':'https://www.amazon.eg/ref=nav_logo','Jumia':'https://www.jumia.com.eg/ar/','Noon':f'https://www.noon.com/egypt-en/search/?q={self.search_word}'}
        url=f'{sites[self.site_name]}'
        with SB(uc=True, test=True, locale="en") as sb:
            sb.sleep(random.randint(1, 5))
            sb.maximize()
            sb.sleep(random.randint(1, 5))
            sb.activate_cdp_mode(url)         
            sb.sleep(random.randint(5, 10))
            try:
                sb.solve_captcha()
                sb.sleep(random.randint(5, 7))
            except:
                print('nocaptcha')  
            if self.site_name =='Amazon':
                sb.type('input[id="twotabsearchtextbox"]',self.search_word)
                self.search_button='input[id="nav-search-submit-button"]'
                sb.click(self.search_button)
                sb.sleep(random.randint(1,5))
                self.html_source = sb.get_page_source()
                self.amazone()
                sb.sleep(random.randint(1,4)) 
                next_button_selector=('a.s-pagination-next')
                while self.search==True:        
                        if sb.is_element_visible(next_button_selector):
                            sb.scroll_into_view(next_button_selector) 
                            sb.sleep(random.randint(2,4))
                            sb.click(next_button_selector)
                            sb.sleep(random.randint(2, 6))
                            self.amazone()
                        else:
                            break
            elif self.site_name =='Jumia':
                sb.click('button.cls')
                sb.sleep(random.randint(1, 3))
                sb.click('Button.-df')
                sb.type('input[name="q"]', self.search_word ,timeout=5)
                sb.sleep(random.randint(5, 10))
                sb.click('Button._search')
                sb.sleep(random.randint(5, 10))
                self.html_source = sb.get_page_source() 
                self.jumia()
                next_button_selector = 'a[aria-label="الصفحة التالية"]'
                while self.search==True:        
                        if sb.is_element_visible(next_button_selector):
                            sb.scroll_into_view(next_button_selector) 
                            sb.sleep(random.randint(2,4))
                            sb.click(next_button_selector)
                            sb.sleep(random.randint(2, 6))
                            self.jumia()
                        else:
                            break

            elif self.site_name =='Noon':
                self.html_source = sb.get_page_source()
                self.noon()
                next_button_selector = 'a[aria-label="Next page"]'
                while self.search==True:        
                        if sb.is_element_visible(next_button_selector):
                            sb.scroll_into_view(next_button_selector) 
                            sb.sleep(random.randint(2,4))
                            sb.click(next_button_selector)
                            sb.sleep(random.randint(2, 6))
                            self.noon()
                        else:
                            break
    def amazone(self) :
        today=str(datetime.date.today())
        file_name=f'{self.search_word}_{self.site_name}_{today}.xlsx'
        titles=["رقم المنتج",'الاسم','السعر','الاتاحة','التقييم']
        if os.path.exists(file_name):
             wb=load_workbook(file_name)
             ws=wb.active
             col_num=ws.max_column
        else:
            wb=Workbook()
            ws=wb.active    
            ws.append(titles)
            col_num=len(titles)
        self.tableWidget.setColumnCount(col_num)
        self.tableWidget.setRowCount(10)
        # 4. Set Header Labels
        self.tableWidget.setHorizontalHeaderLabels(titles)
        QApplication.processEvents()
        self.tableWidget.verticalHeader().setVisible(False)
        pro=[]
        soup = BeautifulSoup(self.html_source, 'html.parser')
        products=soup.find_all('div',{'data-component-type':"s-search-result"})
        i=0
        x=0
        for product in products:
            product_num=str(ws.max_row)
            pro=[]
            pro.append(product_num)
            cont=product.find('div',{'data-cy':'title-recipe'})
            if cont:
                product_link=cont.find('a',{'class':"a-text-normal"})
                h=cont.find('h2')
                if  h!=None :
                    title=h.find('span').text.strip()
                    
                else:
                    title='no title' 
            pro.append(title)
            cont1=product.find('div',{'data-cy':'price-recipe'})
            if  cont1:
                price=cont1.find('span',{'class':"a-price-whole"})
                if price:
                        price1=price.text.strip()
                        pro.append(price1)
                else:
                        pro.append('not found')
            store=product.find('div',{'class':"a-button-stack"})
            add_button=store.find('input',{'name':"submit.addToCart"})
            if add_button:
                    pro.append('available')
            else:
                pro.append('not available')
            evaluate=product.find('div',{'data-cy':'reviews-block'})
            if evaluate:
                 stars=evaluate.find('span').text.strip()
            else:
                 stars='لم يقييم'
            pro.append(stars)
            ws.append(pro)
            r_max=ws.max_row
            pro.append(title)
            if product_link:
                        link=product_link.get('href')
                        total_link=f'https://www.amazon.eg{link}'
            else:
                          total_link='not found'
            ws.cell(row=r_max,column=2).hyperlink=total_link
             
            if (i<=20):
                self.tableWidget.setItem(i, 0, QTableWidgetItem(str(i+1)))
                for j in range (1,len(pro)): 
                    item = QTableWidgetItem(pro[j])
                    self.tableWidget.setItem(i, j, item) 
                    QApplication.processEvents() 
            i+=1
        wb.save(file_name) 
        self.format_file(file_name)              
    def jumia(self):
        today=str(datetime.date.today())
        file_name=f'{self.search_word}_{self.site_name}_{today}.xlsx'
        titles=['product number','Title','السعر (Price)', 'التقييم (Stars)','discount']
        if os.path.exists(file_name):
             wb=load_workbook(file_name)
             ws=wb.active
             col_num=ws.max_column
        else:
            wb=Workbook()
            ws=wb.active
            ws.append(titles)
            col_num=len(titles)
        self.tableWidget.setColumnCount(col_num)
        self.tableWidget.setRowCount(10)
        self.tableWidget.setHorizontalHeaderLabels(titles)
        self.tableWidget.verticalHeader().setVisible(False)
        soup = BeautifulSoup(self.html_source, 'lxml')
        products=soup.find_all('article',{'class':'c-prd'})
        i=0
        for product in products:
            info=[]
            product_num=str(ws.max_row)
            link=product.find('a',{'class':"core"}).get('href')
            total_link=f'https://www.jumia.com.eg{link}'
            data=product.find('div',{'class':'info'})
            title=data.find('h3').text.strip()
            price=data.find('div',{'class':'prc'})
            stars=data.find('div',{'class':'stars'})
            disc=data.find('div',{'class':'_dsct'})
            if title==None :
                title='not found'
            if price != None:
                price=price.text.strip()
            else:
                price='not found'
            if stars != None:
                stars=stars.text.strip()
            else:
                stars='not found'
            if disc != None:
                disc=disc.text.strip()
            else:
                disc='no discount'
            #print(f'product number {i} : Title :{title} , Price is {price} and Number of Stars is {stars}')
            info.append(product_num)
            info.append(title)
            info.append(price)
            info.append(stars)
            info.append(disc)
            ws.append(info)
            r_max=ws.max_row
            ws.cell(row=r_max,column=2).hyperlink=total_link
            if (i<=10):
                self.tableWidget.setItem(i, 0, QTableWidgetItem(str(i+1))) 
                for j in range (1,len(info)):
                  item = QTableWidgetItem(info[j])
                  self.tableWidget.setItem(i, j, item)
                  QApplication.processEvents()
            i+=1
        wb.save(file_name)
        self.format_file(file_name)
    def noon(self):
        today=str(datetime.date.today())
        file_name=f'{self.search_word}_{self.site_name}_{today}.xlsx'
        titles=['product number','Title','السعر (Price)', 'التقييم (Stars)','discount']
        if os.path.exists(file_name):
             wb=load_workbook(file_name)
             ws=wb.active
             col_num=ws.max_column
        else:
            wb=Workbook()
            ws=wb.active
            ws.append(titles)
            col_num=len(titles)
        self.tableWidget.setColumnCount(col_num)
        self.tableWidget.setHorizontalHeaderLabels(titles)
        self.tableWidget.verticalHeader().setVisible(False)
        self.tableWidget.setRowCount(10)
        QApplication.processEvents()
        soup = BeautifulSoup(self.html_source, 'lxml')
        products=soup.find_all('div', {'class':'_linkWrapper_1ts6x_1'})
        i=0
        for product in products:
            info=[]
            link=product.find('a',{'class':"_productBoxLink_1ts6x_5"}).get('href')
            total_link=f'https://www.noon.com{link}'
            product_num=str(ws.max_row)
            info.append(product_num)
            title=product.find('h2',{'class':"_title_i1yaq_19"})
            if title:
                 title_name=title.text.strip()
            else:
                 title_name='Not found'
            info.append(title_name)
            price_div=product.find('div',{'class':"_priceNudgesCtr_i1yaq_30"})
            if price_div:
                price_strong=price_div.find('strong',{'class':"_amount_1o2w0_59"})
                if  price_strong:
                    price=price_strong.text.strip()
            else:
                 price='Not found'
            info.append(price)
            stars_container=product.find('div',{'class':"_container_1r83y_31"})
            if stars_container:
                 stars_div=stars_container.find('div',{'class':"_textCtr_1r83y_16"})
                 if  stars_div:
                     stars=stars_div.text.strip()
            else:
                 stars='not found'
            info.append(stars)
            discount_container=product.find('div',{'class':"_preReductionRow_1o2w0_74"})
            if discount_container:
                 discount_span=discount_container.find('span',{'class':"_discount_grckm_1"})
                 if  discount_span :
                     discount=discount_span.text.strip()
            else:
                 discount='not found'
            info.append(discount)
            ws.append(info)
            r_max=ws.max_row
            ws.cell(row=r_max,column=2).hyperlink=total_link
            if (i<=10):
                self.tableWidget.setItem(i, 0, QTableWidgetItem(str(i+1)))
                for j in range (1,len(info)): 
                    item = QTableWidgetItem(info[j])
                    self.tableWidget.setItem(i, j, item)
                    QApplication.processEvents()
            i+=1
        wb.save(file_name)
        self.format_file(file_name)
    def load_data(self,data):
        for i in range (10) :
             for j in range (len(data)):
                  item = QTableWidgetItem(data[j])
                  self.tableWidget.setItem(i, j, item)
    def open_excel(self):
        today=str(datetime.date.today())
        file_name=f'{self.search_word}_{self.site_name}_{today}.xlsx'
        if os.path.exists(file_name):
            os.startfile(file_name)
        else:
            QMessageBox.warning(self, "خطأ", "ملف Excel غير موجود")              
    def format_file(self, file):
        wb = load_workbook(file)
        ws = wb.active
        bd = Side(style='thin', color="000000")
        for row in ws.iter_rows():
            for cell in row:
                cell.font = Font(bold=True, size=14)
                cell.border = Border(
                    left=bd,
                    right=bd,
                    top=bd,
                    bottom=bd
                )
                cell.alignment = Alignment(
                    horizontal='left',
                    vertical='center'
                )
        for column in ws.columns:
            max_length = 0
            column_letter = get_column_letter(column[0].column)
            for cell in column:
                if cell.value is not None:
                    if len(str(cell.value)) <=50:
                        max_length = max(
                            max_length,
                            len(str(cell.value))
                        )
                    else:
                         max_length=50
            ws.column_dimensions[column_letter].width = max_length + 3
        wb.save(file)
    def stop_searching(self):
         self.search=False
if __name__=="__main__":
    app=QApplication(sys.argv)
    window=scrapping_window()
    window.show()
    sys.exit(app.exec())